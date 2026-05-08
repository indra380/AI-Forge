from __future__ import annotations

import base64
import csv
import json
from dataclasses import dataclass
from pathlib import Path
import re
import uuid
import zipfile

from fastapi import HTTPException, UploadFile
import fitz
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.core.logging import get_logger
from app.models.attachment import Attachment
from app.models.chat_message import ChatMessage
from app.schemas.attachment import AttachmentResponse
from app.services.attachment_storage import AttachmentStorage, get_attachment_storage

logger = get_logger(__name__)

_SAFE_NAME_RE = re.compile(r"[^A-Za-z0-9._-]+")
_CODE_EXTENSIONS = {
    ".py": "text/x-python",
    ".js": "application/javascript",
    ".ts": "application/typescript",
    ".json": "application/json",
    ".yaml": "text/yaml",
    ".yml": "text/yaml",
    ".html": "text/html",
    ".css": "text/css",
    ".java": "text/x-java-source",
    ".cs": "text/x-csharp",
    ".cpp": "text/x-c++src",
    ".sql": "text/x-sql",
}
_TABLE_EXTENSIONS = {
    ".csv": "text/csv",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
_DOCUMENT_EXTENSIONS = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".md": "text/markdown",
}
_IMAGE_EXTENSIONS = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
    ".svg": "image/svg+xml",
}
_VIDEO_EXTENSIONS = {
    ".mp4": "video/mp4",
    ".mov": "video/quicktime",
    ".webm": "video/webm",
}
_FORMULA_EXTENSIONS = {
    ".tex": "application/x-tex",
}


@dataclass(slots=True)
class AttachmentContext:
    attachment: Attachment
    prompt_text: str | None


class AttachmentService:
    def __init__(self, storage: AttachmentStorage):
        self._storage = storage

    async def upload_attachment(
        self,
        db: AsyncSession,
        *,
        user_id: uuid.UUID,
        session_id: uuid.UUID,
        upload: UploadFile,
    ) -> Attachment:
        original_name = Path(upload.filename or "upload.bin").name
        safe_stem = self._sanitize_stem(Path(original_name).stem)
        suffix = Path(original_name).suffix.lower()
        attachment_type, allowed_content_type = self._resolve_allowed_type(suffix)

        content_type = (upload.content_type or allowed_content_type or "application/octet-stream").lower()
        if suffix == ".json":
            attachment_type, allowed_content_type = await self._resolve_json_attachment_type(upload, content_type)
        if not self._is_content_type_allowed(attachment_type, content_type, allowed_content_type):
            raise HTTPException(
                status_code=415,
                detail={"error": "unsupported_file_type", "message": "Unsupported file type"},
            )

        stored = await self._storage.save(upload, self._category_dir(attachment_type), safe_stem)
        max_size_bytes = settings.MAX_UPLOAD_SIZE_MB * 1024 * 1024
        if stored.file_size > max_size_bytes:
            Path(stored.storage_path).unlink(missing_ok=True)
            raise HTTPException(
                status_code=413,
                detail={"error": "file_too_large", "message": "Attachment exceeds upload size limit"},
            )

        await self._run_antivirus_scan_hook(stored.storage_path)
        extracted_text = self._extract_text(Path(stored.storage_path), attachment_type, suffix)

        attachment = Attachment(
            user_id=user_id,
            session_id=session_id,
            message_id=None,
            file_name=stored.file_name,
            original_name=original_name,
            content_type=content_type,
            file_size=stored.file_size,
            storage_path=stored.storage_path,
            attachment_type=attachment_type,
            extracted_text=extracted_text,
        )
        db.add(attachment)
        await db.commit()
        await db.refresh(attachment)
        return attachment

    async def _resolve_json_attachment_type(
        self, upload: UploadFile, content_type: str
    ) -> tuple[str, str]:
        payload = await upload.read(256 * 1024)
        await upload.seek(0)
        try:
            parsed = json.loads(payload.decode("utf-8", errors="ignore"))
            if isinstance(parsed, list) and all(isinstance(item, dict) for item in parsed[:20]):
                return "table", "application/json"
        except Exception:
            pass

        # Default .json to code unless it clearly matches tabular JSON.
        return "code", "application/json"

    async def list_message_attachments(
        self, db: AsyncSession, messages: list[ChatMessage]
    ) -> dict[uuid.UUID, list[AttachmentResponse]]:
        if not messages:
            return {}

        message_ids = [message.id for message in messages]
        stmt = select(Attachment).where(Attachment.message_id.in_(message_ids)).order_by(Attachment.uploaded_at.asc())
        result = await db.execute(stmt)
        attachments = list(result.scalars().all())

        mapped: dict[uuid.UUID, list[AttachmentResponse]] = {message_id: [] for message_id in message_ids}
        for attachment in attachments:
            if attachment.message_id is None:
                continue
            mapped.setdefault(attachment.message_id, []).append(self.to_response(attachment))
        return mapped

    async def attach_uploaded_files_to_message(
        self,
        db: AsyncSession,
        *,
        user_id: uuid.UUID,
        session_id: uuid.UUID,
        message_id: uuid.UUID,
        attachment_ids: list[uuid.UUID],
    ) -> list[Attachment]:
        if not attachment_ids:
            return []

        stmt = select(Attachment).where(
            Attachment.id.in_(attachment_ids),
            Attachment.user_id == user_id,
            Attachment.session_id == session_id,
        )
        result = await db.execute(stmt)
        attachments = list(result.scalars().all())
        if len(attachments) != len(set(attachment_ids)):
            raise HTTPException(
                status_code=404,
                detail={"error": "attachment_not_found", "message": "One or more attachments are unavailable"},
            )

        for attachment in attachments:
            attachment.message_id = message_id

        await db.commit()
        for attachment in attachments:
            await db.refresh(attachment)
        return attachments

    async def build_attachment_context(
        self,
        db: AsyncSession,
        *,
        user_id: uuid.UUID,
        session_id: uuid.UUID,
        attachment_ids: list[uuid.UUID],
    ) -> str:
        if not attachment_ids:
            return ""

        stmt = select(Attachment).where(
            Attachment.id.in_(attachment_ids),
            Attachment.user_id == user_id,
            Attachment.session_id == session_id,
        )
        result = await db.execute(stmt)
        attachments = list(result.scalars().all())
        if not attachments:
            return ""

        parts: list[str] = []
        for attachment in attachments:
            if attachment.attachment_type == "image":
                parts.append(
                    f"Attachment: {attachment.original_name}\nAttachment context: image attachment uploaded. Analyze the visual content."
                )
                continue
            if attachment.attachment_type == "video":
                parts.append(
                    f"Attachment: {attachment.original_name}\nAttachment context: video attachment uploaded. Analyze its visual and audio content."
                )
                continue
            if attachment.extracted_text:
                parts.append(
                    f"Attachment: {attachment.original_name}\nAttachment context: the following text was extracted from the uploaded file and should be treated as the file's readable content.\nExtracted content:\n{attachment.extracted_text[:6000]}"
                )
            else:
                parts.append(
                    f"Attachment: {attachment.original_name}\nAttachment context: file uploaded with content type {attachment.content_type}, but no readable text could be extracted."
                )

        return "\n\n".join(parts)

    async def build_multimodal_user_content(
        self,
        db: AsyncSession,
        *,
        user_id: uuid.UUID,
        session_id: uuid.UUID,
        attachment_ids: list[uuid.UUID],
        message_text: str,
    ) -> str | list[dict[str, object]]:
        if not attachment_ids:
            return message_text

        stmt = select(Attachment).where(
            Attachment.id.in_(attachment_ids),
            Attachment.user_id == user_id,
            Attachment.session_id == session_id,
        )
        result = await db.execute(stmt)
        attachments = list(result.scalars().all())
        if not attachments:
            return message_text

        content_blocks: list[dict[str, object]] = []
        text_parts: list[str] = []
        if message_text.strip():
            text_parts.append(message_text.strip())

        for attachment in attachments:
            if attachment.attachment_type == "image":
                text_parts.append(
                    f"Analyze the uploaded image named {attachment.original_name}."
                )
                data_url = self._to_data_url(Path(attachment.storage_path), attachment.content_type)
                content_blocks.append(
                    {"type": "image_url", "image_url": {"url": data_url}}
                )
                continue

            if attachment.attachment_type == "video":
                video_path = Path(attachment.storage_path)
                max_video_bytes = 50 * 1024 * 1024  # 50 MB inline limit
                if video_path.stat().st_size <= max_video_bytes:
                    text_parts.append(
                        f"Analyze the uploaded video named {attachment.original_name}. Describe what happens visually and summarize any spoken content."
                    )
                    data_url = self._to_data_url(video_path, attachment.content_type)
                    content_blocks.append(
                        {"type": "image_url", "image_url": {"url": data_url}}
                    )
                else:
                    text_parts.append(
                        f"The video file {attachment.original_name} ({attachment.content_type}) is too large to process inline. Only metadata is available."
                    )
                continue

            if (
                attachment.attachment_type == "document"
                and Path(attachment.original_name).suffix.lower() == ".pdf"
                and not attachment.extracted_text
            ):
                text_parts.append(
                    f"The uploaded PDF named {attachment.original_name} had no direct text extraction. Read the rendered page images instead."
                )
                for page_data_url in self._render_pdf_pages_to_data_urls(Path(attachment.storage_path)):
                    content_blocks.append(
                        {"type": "image_url", "image_url": {"url": page_data_url}}
                    )

        if not content_blocks:
            return message_text

        combined_text = "\n\n".join(part for part in text_parts if part).strip()
        if combined_text:
            return [{"type": "text", "text": combined_text}, *content_blocks]
        return content_blocks

    async def delete_pending_attachment(
        self,
        db: AsyncSession,
        *,
        user_id: uuid.UUID,
        session_id: uuid.UUID,
        attachment_id: uuid.UUID,
    ) -> None:
        stmt = select(Attachment).where(
            Attachment.id == attachment_id,
            Attachment.user_id == user_id,
            Attachment.session_id == session_id,
        )
        result = await db.execute(stmt)
        attachment = result.scalar_one_or_none()
        if attachment is None:
            raise HTTPException(
                status_code=404,
                detail={"error": "attachment_not_found", "message": "Attachment not found"},
            )
        if attachment.message_id is not None:
            raise HTTPException(
                status_code=409,
                detail={"error": "attachment_locked", "message": "Attachment already sent"},
            )

        Path(attachment.storage_path).unlink(missing_ok=True)
        await db.delete(attachment)
        await db.commit()

    def to_response(self, attachment: Attachment) -> AttachmentResponse:
        file_url = self._to_public_url(attachment.storage_path)
        return AttachmentResponse(
            id=attachment.id,
            file_id=attachment.id,
            file_name=attachment.file_name,
            original_name=attachment.original_name,
            file_url=file_url,
            content_type=attachment.content_type,
            file_size=attachment.file_size,
            attachment_type=attachment.attachment_type,
            extracted_text=attachment.extracted_text,
            uploaded_at=attachment.uploaded_at,
        )

    @staticmethod
    def _sanitize_stem(stem: str) -> str:
        cleaned = _SAFE_NAME_RE.sub("-", stem).strip("-._")
        return cleaned[:80] or "attachment"

    @staticmethod
    def _category_dir(attachment_type: str) -> str:
        return {
            "image": "images",
            "video": "videos",
            "table": "tables",
            "document": "documents",
            "code": "code",
            "formula": "formulas",
        }[attachment_type]

    def _resolve_allowed_type(self, suffix: str) -> tuple[str, str]:
        if suffix in _IMAGE_EXTENSIONS:
            return "image", _IMAGE_EXTENSIONS[suffix]
        if suffix in _VIDEO_EXTENSIONS:
            return "video", _VIDEO_EXTENSIONS[suffix]
        if suffix in _TABLE_EXTENSIONS:
            return "table", _TABLE_EXTENSIONS[suffix]
        if suffix in _FORMULA_EXTENSIONS:
            return "formula", _FORMULA_EXTENSIONS[suffix]
        if suffix in _DOCUMENT_EXTENSIONS:
            return "document", _DOCUMENT_EXTENSIONS[suffix]
        if suffix in _CODE_EXTENSIONS:
            return "code", _CODE_EXTENSIONS[suffix]
        raise HTTPException(
            status_code=415,
            detail={"error": "unsupported_file_type", "message": "Unsupported file extension"},
        )

    def _is_content_type_allowed(
        self, attachment_type: str, content_type: str, expected_content_type: str
    ) -> bool:
        allowed = {
            "image": settings.ATTACHMENT_IMAGE_MIME_TYPES,
            "video": settings.ATTACHMENT_VIDEO_MIME_TYPES,
            "table": settings.ATTACHMENT_TABLE_MIME_TYPES,
            "document": settings.ATTACHMENT_DOCUMENT_MIME_TYPES,
            "code": settings.ATTACHMENT_CODE_MIME_TYPES,
            "formula": settings.ATTACHMENT_FORMULA_MIME_TYPES,
        }[attachment_type]
        return content_type in {mime.lower() for mime in allowed} or content_type == expected_content_type

    async def _run_antivirus_scan_hook(self, storage_path: str) -> None:
        logger.debug("Antivirus hook placeholder for %s", storage_path)

    def _extract_text(self, path: Path, attachment_type: str, suffix: str) -> str | None:
        try:
            if attachment_type == "code":
                return path.read_text(encoding="utf-8", errors="ignore")
            if attachment_type == "formula":
                return path.read_text(encoding="utf-8", errors="ignore")
            if attachment_type == "table":
                return self._extract_table_text(path, suffix)
            if suffix == ".txt":
                return path.read_text(encoding="utf-8", errors="ignore")
            if suffix == ".md":
                return path.read_text(encoding="utf-8", errors="ignore")
            if suffix == ".pdf":
                reader = PdfReader(str(path))
                return "\n".join(page.extract_text() or "" for page in reader.pages).strip() or None
            if suffix == ".docx":
                with zipfile.ZipFile(path) as archive:
                    content = archive.read("word/document.xml").decode("utf-8", errors="ignore")
                text = re.sub(r"<[^>]+>", " ", content)
                return re.sub(r"\s+", " ", text).strip() or None
        except Exception:
            logger.warning("Attachment text extraction failed for %s", path)
        return None

    @staticmethod
    def _extract_table_text(path: Path, suffix: str) -> str | None:
        if suffix == ".csv":
            with path.open("r", encoding="utf-8", errors="ignore", newline="") as handle:
                rows = list(csv.reader(handle))
            return AttachmentService._rows_to_markdown_table(rows)

        if suffix == ".xlsx":
            workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
            try:
                sheet = workbook.active
                rows: list[list[str]] = []
                for row in sheet.iter_rows(values_only=True):
                    rows.append(["" if cell is None else str(cell) for cell in row])
                return AttachmentService._rows_to_markdown_table(rows)
            finally:
                workbook.close()

        # JSON tables are accepted when uploaded with .json extension.
        if suffix == ".json":
            raw = path.read_text(encoding="utf-8", errors="ignore")
            parsed = json.loads(raw)
            return AttachmentService._json_to_markdown_table(parsed)

        return None

    @staticmethod
    def _rows_to_markdown_table(rows: list[list[str]], max_rows: int = 40) -> str | None:
        if not rows:
            return None

        normalized = [[str(cell).strip() for cell in row] for row in rows if row]
        if not normalized:
            return None

        width = max(len(row) for row in normalized)
        padded = [row + [""] * (width - len(row)) for row in normalized[: max_rows + 1]]
        header = padded[0]
        body = padded[1:]

        sep = ["---"] * width
        lines = [
            "| " + " | ".join(header) + " |",
            "| " + " | ".join(sep) + " |",
        ]
        lines.extend("| " + " | ".join(row) + " |" for row in body)

        if len(normalized) > max_rows + 1:
            lines.append(f"\n... truncated {len(normalized) - (max_rows + 1)} rows ...")

        return "\n".join(lines)

    @staticmethod
    def _json_to_markdown_table(data: object) -> str:
        if isinstance(data, list) and data and all(isinstance(item, dict) for item in data):
            keys: list[str] = []
            for item in data:
                for key in item.keys():
                    if key not in keys:
                        keys.append(key)
            rows = [keys]
            for item in data[:40]:
                rows.append([str(item.get(key, "")) for key in keys])
            return AttachmentService._rows_to_markdown_table(rows) or "[]"

        return json.dumps(data, indent=2)

    def _to_public_url(self, storage_path: str) -> str:
        upload_root = Path(settings.UPLOAD_PATH).resolve()
        file_path = Path(storage_path).resolve()
        relative = file_path.relative_to(upload_root).as_posix()
        return f"/uploads/{relative}"

    @staticmethod
    def _to_data_url(path: Path, content_type: str) -> str:
        encoded = base64.b64encode(path.read_bytes()).decode("ascii")
        return f"data:{content_type};base64,{encoded}"

    @staticmethod
    def _render_pdf_pages_to_data_urls(path: Path, max_pages: int = 4) -> list[str]:
        data_urls: list[str] = []
        document = fitz.open(path)
        try:
            for page_index in range(min(len(document), max_pages)):
                page = document.load_page(page_index)
                pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                encoded = base64.b64encode(pixmap.tobytes("png")).decode("ascii")
                data_urls.append(f"data:image/png;base64,{encoded}")
        finally:
            document.close()
        return data_urls


_attachment_service_singleton: AttachmentService | None = None


def get_attachment_service() -> AttachmentService:
    global _attachment_service_singleton
    if _attachment_service_singleton is None:
        _attachment_service_singleton = AttachmentService(get_attachment_storage())
    return _attachment_service_singleton